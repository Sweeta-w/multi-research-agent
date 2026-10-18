import io
import os

import streamlit as st
from deep_translator import GoogleTranslator
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from crew.research_crew import AGENT_ORDER, build_crew

# ============================================================
# HELPER FUNCTIONS FOR TRANSLATION & EXPORTS
# ============================================================


def translate_text(text: str, target_lang: str) -> str:
    """Translates research report text to Urdu, Hindi, or Sindhi."""
    lang_codes = {
        "Urdu": "ur",
        "Hindi": "hi",
        "Sindhi": "sd",
    }
    if target_lang in lang_codes:
        try:
            translator = GoogleTranslator(
                source="en", target=lang_codes[target_lang]
            )
            chunks = [text[i : i + 4000] for i in range(0, len(text), 4000)]
            translated_chunks = [translator.translate(chunk) for chunk in chunks]
            return "\n\n".join(translated_chunks)
        except Exception as e:
            st.error(f"Translation Error: {e}")
            return text
    return text


def create_docx(text: str, title: str = "Research Report") -> io.BytesIO:
    """Generates a Word document in memory."""
    doc = Document()
    doc.add_heading(title, level=1)
    for paragraph in text.split("\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph)
    bio = io.BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


def create_pdf(text: str, title: str = "Research Report") -> io.BytesIO:
    """Generates a PDF document in memory."""
    bio = io.BytesIO()
    pdf = SimpleDocTemplate(bio, pagesize=letter)
    styles = getSampleStyleSheet()
    normal_style = styles["Normal"]
    normal_style.leading = 14

    story = [
        Paragraph(f"<b>{title}</b>", styles["Title"]),
        Spacer(1, 12),
    ]
    for line in text.split("\n"):
        if line.strip():
            clean_line = (
                line.replace("<", "&lt;").replace(">", "&gt;").replace("&", "&amp;")
            )
            story.append(Paragraph(clean_line, normal_style))
            story.append(Spacer(1, 6))

    pdf.build(story)
    bio.seek(0)
    return bio


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="ResearchFlow AI",
    page_icon="🔎",
    layout="wide",
    initial_sidebar_state="collapsed",
)


# ============================================================
# CUSTOM CSS
# ============================================================

st.html(
    """
    <style>

    /* =====================================================
       GLOBAL & TEXT FIXES
    ===================================================== */

    .stApp {
        background:
            radial-gradient(
                circle at 10% 10%,
                rgba(99, 102, 241, 0.10),
                transparent 30%
            ),
            radial-gradient(
                circle at 90% 20%,
                rgba(14, 165, 233, 0.08),
                transparent 30%
            ),
            #f8fafc;
        color: #0f172a !important;
    }

    .stApp p, .stApp span, .stApp label, .stApp h1, .stApp h2, .stApp h3 {
        color: #0f172a;
    }

    .block-container {
        max-width: 1180px;
        padding-top: 2rem;
        padding-bottom: 4rem;
    }


    /* =====================================================
       INPUT AREA & SELECTBOX FIX
    ===================================================== */

    .stTextArea textarea {
        background-color: #ffffff !important;
        color: #0f172a !important;
        border: 1px solid #cbd5e1 !important;
        border-radius: 12px !important;
        font-size: 0.95rem !important;
    }

    .stTextArea textarea::placeholder {
        color: #94a3b8 !important;
    }

    .stTextArea textarea:focus {
        border-color: #6366f1 !important;
        box-shadow: 0 0 0 3px rgba(99, 102, 241, 0.15) !important;
    }

    div[data-baseweb="select"] > div {
        background-color: #ffffff !important;
        border-radius: 12px !important;
        border: 1px solid #cbd5e1 !important;
    }


    /* =====================================================
       HERO
    ===================================================== */

    .hero {
        padding: 2.5rem 1rem 2rem 1rem;
        text-align: center;
    }

    .hero-badge {
        display: inline-block;
        padding: 0.45rem 0.85rem;
        border-radius: 999px;
        background: rgba(99, 102, 241, 0.10);
        border: 1px solid rgba(99, 102, 241, 0.15);
        color: #4f46e5 !important;
        font-size: 0.78rem;
        font-weight: 700;
        letter-spacing: 0.04em;
        margin-bottom: 1rem;
    }

    .hero-title {
        margin: 0;
        font-size: clamp(
            2.3rem,
            6vw,
            4.5rem
        );
        line-height: 1.04;
        font-weight: 800;
        letter-spacing: -0.05em;
        color: #0f172a !important;
    }

    .hero-title span {
        color: #4f46e5 !important;
    }

    .hero-subtitle {
        max-width: 720px;
        margin: 1.25rem auto 0 auto;
        color: #64748b !important;
        font-size: clamp(
            0.95rem,
            2vw,
            1.12rem
        );
        line-height: 1.7;
    }


    /* =====================================================
       CARDS & SECTION TEXT
    ===================================================== */

    .glass-card {
        background: rgba(255, 255, 255, 0.82);
        border: 1px solid rgba(226, 232, 240, 0.95);
        border-radius: 22px;
        padding: 1.4rem;
        box-shadow: 0 10px 35px rgba(15, 23, 42, 0.06);
        backdrop-filter: blur(12px);
        margin-bottom: 1rem;
    }

    .section-title {
        font-size: 1.1rem;
        font-weight: 750;
        color: #0f172a !important;
        margin-bottom: 0.3rem;
    }

    .section-description {
        color: #64748b !important;
        font-size: 0.88rem;
        line-height: 1.5;
        margin-bottom: 0;
    }


    /* =====================================================
       AGENT CARDS
    ===================================================== */

    .agent-card {
        display: flex;
        align-items: center;
        gap: 0.9rem;
        padding: 0.9rem;
        border-radius: 15px;
        border: 1px solid #e2e8f0;
        background: #ffffff;
        margin-bottom: 0.7rem;
        transition: transform 0.2s ease, box-shadow 0.2s ease;
    }

    .agent-card:hover {
        transform: translateY(-1px);
        box-shadow: 0 8px 20px rgba(15, 23, 42, 0.06);
    }

    .agent-icon {
        width: 42px;
        height: 42px;
        min-width: 42px;
        display: flex;
        align-items: center;
        justify-content: center;
        border-radius: 12px;
        background: #eef2ff;
        font-size: 1.15rem;
        color: #4f46e5;
    }

    .agent-name {
        font-weight: 700;
        color: #0f172a !important;
        font-size: 0.92rem;
    }

    .agent-description {
        color: #64748b !important;
        font-size: 0.78rem;
        line-height: 1.4;
        margin-top: 0.15rem;
    }

    .working {
        border-color: #818cf8;
        background: linear-gradient(135deg, #eef2ff, #ffffff);
        box-shadow: 0 0 0 1px rgba(99, 102, 241, 0.08), 0 10px 25px rgba(99, 102, 241, 0.08);
    }

    .completed {
        border-color: #bbf7d0;
        background: linear-gradient(135deg, #f0fdf4, #ffffff);
    }

    .waiting {
        opacity: 0.58;
    }

    .current-agent {
        display: flex;
        align-items: center;
        gap: 0.8rem;
        padding: 1rem 1.1rem;
        border-radius: 16px;
        border: 1px solid rgba(129, 140, 248, 0.35);
        background: linear-gradient(135deg, #eef2ff, #ffffff);
        margin-bottom: 1rem;
    }

    .current-agent-dot {
        width: 11px;
        height: 11px;
        min-width: 11px;
        border-radius: 50%;
        background: #6366f1;
        box-shadow: 0 0 0 5px rgba(99, 102, 241, 0.12);
    }

    .current-agent-title {
        color: #312e81 !important;
        font-size: 0.95rem;
        font-weight: 750;
    }

    .current-agent-description {
        color: #64748b !important;
        font-size: 0.8rem;
        margin-top: 0.2rem;
    }


    /* =====================================================
       REPORT CONTAINER
    ===================================================== */

    .report-container {
        background: #ffffff;
        border: 1px solid #e2e8f0;
        border-radius: 22px;
        padding: clamp(1.2rem, 4vw, 2.5rem);
        box-shadow: 0 15px 45px rgba(15, 23, 42, 0.06);
        line-height: 1.75;
        color: #0f172a !important;
        margin-top: 1rem;
    }

    .report-container p, 
    .report-container li, 
    .report-container td, 
    .report-container th, 
    .report-container span {
        color: #334155 !important;
    }

    .report-container h1, 
    .report-container h2, 
    .report-container h3, 
    .report-container h4 {
        color: #0f172a !important;
        font-weight: 750;
    }

    .report-container table {
        width: 100%;
        border-collapse: collapse;
        margin: 1rem 0;
    }

    .report-container th, .report-container td {
        border: 1px solid #e2e8f0;
        padding: 0.75rem;
        text-align: left;
    }

    .report-container th {
        background-color: #f8fafc;
    }

    .footer {
        text-align: center;
        color: #94a3b8 !important;
        font-size: 0.78rem;
        padding-top: 2.5rem;
    }

    @media (max-width: 768px) {
        .block-container {
            padding-left: 1rem;
            padding-right: 1rem;
            padding-top: 1rem;
        }

        .hero {
            padding: 1.5rem 0.5rem 1.5rem;
        }

        .hero-title {
            font-size: 2.5rem;
        }

        .hero-subtitle {
            font-size: 0.95rem;
        }

        .glass-card, .report-container {
            border-radius: 16px;
            padding: 1rem;
        }

        .agent-card {
            padding: 0.75rem;
        }

        .agent-icon {
            width: 38px;
            height: 38px;
            min-width: 38px;
        }
    }

    </style>
    """
)


# ============================================================
# HERO
# ============================================================

st.html(
    """
    <div class="hero">
        <div class="hero-badge">✦ MULTI-AGENT RESEARCH SYSTEM</div>
        <h1 class="hero-title">
            Research smarter with<br>
            <span>AI teammates.</span>
        </h1>
        <p class="hero-subtitle">
            A CrewAI-powered research team that plans your question,
            searches for evidence, checks the findings, and writes
            a structured research report.
        </p>
    </div>
    """
)


# ============================================================
# MAIN LAYOUT
# ============================================================

left_column, right_column = st.columns([1.55, 1], gap="large")


# ============================================================
# LEFT — RESEARCH INPUT
# ============================================================

with left_column:

    st.html(
        """
        <div class="glass-card">
            <div class="section-title">What do you want to research?</div>
            <div class="section-description">
                Ask a focused question. Your AI research team will investigate it step by step.
            </div>
        </div>
        """
    )

    topic = st.text_area(
        "Research question",
        placeholder="Example: How is generative AI changing software development?",
        height=150,
        label_visibility="collapsed",
    )

    start_research = st.button(
        "🚀  Start Research",
        type="primary",
        use_container_width=True,
    )


# ============================================================
# RIGHT — RESEARCH TEAM
# ============================================================

with right_column:

    st.html(
        """
        <div class="glass-card">
            <div class="section-title">Your research team</div>
            <div class="section-description">Four specialized agents work sequentially.</div>
        </div>
        """
    )

    for agent in AGENT_ORDER:

        st.html(
            f"""
            <div class="agent-card">
                <div class="agent-icon">{agent["icon"]}</div>
                <div>
                    <div class="agent-name">{agent["name"]}</div>
                    <div class="agent-description">{agent["description"]}</div>
                </div>
            </div>
            """
        )


# ============================================================
# START RESEARCH
# ============================================================

if start_research:

    if not topic.strip():
        st.warning("Please enter a research question first.")
        st.stop()

    if not os.getenv("GROQ_API_KEY"):
        st.error("GROQ_API_KEY is missing. Add it in Streamlit Secrets.")
        st.stop()

    if not os.getenv("SERPER_API_KEY"):
        st.error("SERPER_API_KEY is missing. Add it in Streamlit Secrets.")
        st.stop()

    # Reset previous report states
    if "english_report" in st.session_state:
        del st.session_state["english_report"]
    if "current_report" in st.session_state:
        del st.session_state["current_report"]
    if "current_lang" in st.session_state:
        del st.session_state["current_lang"]

    # ========================================================
    # PROGRESS UI
    # ========================================================

    st.divider()

    st.subheader("Research progress")

    current_status = st.empty()
    progress_bar = st.progress(0)
    status_cards = [st.empty() for _ in AGENT_ORDER]

    # ========================================================
    # STATUS CALLBACK
    # ========================================================

    def update_status(completed_agent, next_agent):

        completed_lower = str(completed_agent).lower()

        if "planner" in completed_lower:
            completed_index = 0
            next_index = 1
        elif "research" in completed_lower:
            completed_index = 1
            next_index = 2
        elif "fact" in completed_lower:
            completed_index = 2
            next_index = 3
        else:
            completed_index = 3
            next_index = 4

        completed_data = AGENT_ORDER[completed_index]

        status_cards[completed_index].html(
            f"""
            <div class="agent-card completed">
                <div class="agent-icon" style="background:#dcfce7; color:#16a34a;">✓</div>
                <div>
                    <div class="agent-name">{completed_data["name"]} — Completed</div>
                    <div class="agent-description">{completed_data["description"]}</div>
                </div>
            </div>
            """
        )

        if next_index < len(AGENT_ORDER):

            next_data = AGENT_ORDER[next_index]

            current_status.html(
                f"""
                <div class="current-agent">
                    <div class="current-agent-dot"></div>
                    <div>
                        <div class="current-agent-title">
                            {next_data["icon"]} Currently working: {next_data["name"]}
                        </div>
                        <div class="current-agent-description">
                            {next_data["description"]}
                        </div>
                    </div>
                </div>
                """
            )

            status_cards[next_index].html(
                f"""
                <div class="agent-card working">
                    <div class="agent-icon">{next_data["icon"]}</div>
                    <div>
                        <div class="agent-name">{next_data["name"]} — Working</div>
                        <div class="agent-description">{next_data["description"]}</div>
                    </div>
                </div>
                """
            )

            progress_bar.progress(int(next_index / len(AGENT_ORDER) * 100))

    # ========================================================
    # INITIAL STATE
    # ========================================================

    first_agent = AGENT_ORDER[0]

    current_status.html(
        f"""
        <div class="current-agent">
            <div class="current-agent-dot"></div>
            <div>
                <div class="current-agent-title">
                    {first_agent["icon"]} Currently working: {first_agent["name"]}
                </div>
                <div class="current-agent-description">{first_agent["description"]}</div>
            </div>
        </div>
        """
    )

    status_cards[0].html(
        f"""
        <div class="agent-card working">
            <div class="agent-icon">{first_agent["icon"]}</div>
            <div>
                <div class="agent-name">{first_agent["name"]} — Working</div>
                <div class="agent-description">{first_agent["description"]}</div>
            </div>
        </div>
        """
    )

    for index in range(1, len(AGENT_ORDER)):
        agent = AGENT_ORDER[index]
        status_cards[index].html(
            f"""
            <div class="agent-card waiting">
                <div class="agent-icon">{agent["icon"]}</div>
                <div>
                    <div class="agent-name">{agent["name"]}</div>
                    <div class="agent-description">Waiting for previous agent.</div>
                </div>
            </div>
            """
        )

    # ========================================================
    # RUN CREW
    # ========================================================

    try:

        research_crew = build_crew(status_callback=update_status)

        result = research_crew.kickoff(inputs={"topic": topic.strip()})

        progress_bar.progress(100)

        current_status.success("✅ All research agents completed.")

        raw_report = result.raw if hasattr(result, "raw") else str(result)

        # Store initial English report in session
        st.session_state["english_report"] = raw_report
        st.session_state["current_report"] = raw_report
        st.session_state["current_lang"] = "English"

    except Exception as error:

        current_status.error("Research stopped because an error occurred.")
        st.error(f"Error: {error}")


# ============================================================
# REPORT DISPLAY, TRANSLATION & EXPORTS
# ============================================================

if "english_report" in st.session_state:

    st.divider()
    st.subheader("Research Report")

    # Translation & Controls Block
    with st.expander("🌐 Translate Report & Download Settings", expanded=True):
        col_lang, col_btn = st.columns([2, 1])

        with col_lang:
            target_lang = st.selectbox(
                "Convert Report Language:",
                options=["English", "Urdu", "Hindi", "Sindhi"],
                index=["English", "Urdu", "Hindi", "Sindhi"].index(
                    st.session_state.get("current_lang", "English")
                ),
            )

        with col_btn:
            st.write(" ")
            st.write(" ")
            if st.button("Translate Report", use_container_width=True):
                if target_lang == "English":
                    st.session_state["current_report"] = st.session_state["english_report"]
                    st.session_state["current_lang"] = "English"
                else:
                    with st.spinner(f"Translating report into {target_lang}..."):
                        translated = translate_text(
                            st.session_state["english_report"], target_lang
                        )
                        st.session_state["current_report"] = translated
                        st.session_state["current_lang"] = target_lang
                st.rerun()

    # Download Buttons (Uses currently active translated/original report)
    col_dl1, col_dl2, _ = st.columns([1, 1, 2])
    lang_tag = st.session_state.get("current_lang", "English")

    with col_dl1:
        doc_file = create_docx(
            st.session_state["current_report"],
            title=f"Research Report ({lang_tag})",
        )
        st.download_button(
            label=f"📄 Download Word ({lang_tag})",
            data=doc_file,
            file_name=f"Research_Report_{lang_tag}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True,
        )

    with col_dl2:
        pdf_file = create_pdf(
            st.session_state["current_report"],
            title=f"Research Report ({lang_tag})",
        )
        st.download_button(
            label=f"📕 Download PDF ({lang_tag})",
            data=pdf_file,
            file_name=f"Research_Report_{lang_tag}.pdf",
            mime="application/pdf",
            use_container_width=True,
        )

    # Display Active Report
    st.markdown(
        f'<div class="report-container">\n\n{st.session_state["current_report"]}\n\n</div>',
        unsafe_allow_html=True,
    )


# ============================================================
# FOOTER
# ============================================================

st.html(
    """
    <div class="footer">
        ResearchFlow AI · By Sweeta · Evidence-driven multi-agent research
    </div>
    """
)
